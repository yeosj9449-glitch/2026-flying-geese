"""대시보드 엑셀 공용 서식 헬퍼 (헤더/강조색/열너비)."""
from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="2F5D50")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FILL_GOOD = PatternFill("solid", fgColor="D9EAD3")
FILL_WARN = PatternFill("solid", fgColor="FFF2CC")
FILL_BAD = PatternFill("solid", fgColor="F4CCCC")
FILL_KIDS = PatternFill("solid", fgColor="D0E0F7")


def write_title(sheet: Worksheet, title: str, subtitle: str = "") -> int:
    """시트 상단 제목을 쓰고, 다음에 쓸 행 번호를 돌려준다."""
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=14)
    if subtitle:
        sheet["A2"] = subtitle
        sheet["A2"].font = Font(italic=True, color="666666")
        return 4
    return 3


def write_table(
    sheet: Worksheet,
    start_row: int,
    headers: list[str],
    rows: list[list],
    widths: list[int] | None = None,
) -> int:
    """헤더 + 행을 쓰고 마지막으로 쓴 행 번호를 돌려준다."""
    for col_idx, header in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_idx = start_row
    for row_idx, values in enumerate(rows, start=start_row + 1):
        for col_idx, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(value, str) and len(value) > 20)

    if widths:
        for col_idx, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(col_idx)].width = width
    if rows:
        sheet.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{row_idx}"
    sheet.freeze_panes = sheet.cell(row=start_row + 1, column=1)
    return row_idx


def fill_row(sheet: Worksheet, row_idx: int, n_cols: int, fill: PatternFill) -> None:
    for col_idx in range(1, n_cols + 1):
        sheet.cell(row=row_idx, column=col_idx).fill = fill
