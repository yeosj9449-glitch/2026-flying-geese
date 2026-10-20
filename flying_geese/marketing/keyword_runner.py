"""키워드 발굴 실행 + 리포트 엑셀/업로드용 CSV 생성 (CLI `keywords`)."""
from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from flying_geese.config import Settings, load_settings
from flying_geese.marketing.ad_performance import BidRecommendation, BidRules, load_performance, recommend_all
from flying_geese.marketing.keyword_discovery import (
    DEFAULT_MARKETING_DIR,
    GRADE_CONTENT,
    GRADE_CORE,
    GRADE_EXCLUDE,
    GRADE_ORDER,
    GRADE_TEST,
    ClassifiedKeyword,
    KeywordRules,
    build_hint_keywords,
    classify_all,
    fetch_stats_from_api,
    load_stats_file,
    saturation_label,
    update_keyword_db,
)
from flying_geese.tabular import write_csv_rows
from flying_geese.xlsx_style import FILL_GOOD, FILL_WARN, write_table, write_title

OUTPUT_DIR = Path(__file__).resolve().parents[2] / "output"
SHOPPING_LOOKUP_LIMIT = 150  # 쇼핑 검색 API 일 호출 한도 보호


@dataclass
class KeywordRunResult:
    keywords: list[ClassifiedKeyword]
    source: str
    report_path: Path
    upload_path: Path
    bid_recommendations: list[BidRecommendation] = field(default_factory=list)
    db_path: Path | None = None
    using_examples: bool = False


def _first_existing(*paths: Path) -> Path | None:
    return next((p for p in paths if p.exists()), None)


def run_keyword_discovery(
    marketing_dir: Path = DEFAULT_MARKETING_DIR,
    stats_file: Path | None = None,
    items: list[str] | None = None,
    with_shopping: bool = False,
    settings: Settings | None = None,
    output_dir: Path = OUTPUT_DIR,
    today: date | None = None,
    searchad_client=None,
    shopping_client=None,
) -> KeywordRunResult:
    today = today or date.today()
    settings = settings or load_settings()
    rules = KeywordRules.load(marketing_dir / "keyword_rules.json")
    hints = build_hint_keywords(rules, items)
    using_examples = False

    # 1) 키워드 통계 수집: 지정 파일 → 운영자 파일 → API → 예시 파일 순
    local_file = stats_file or _first_existing(
        marketing_dir / "keyword_stats.csv", marketing_dir / "keyword_stats.json"
    )
    if local_file:
        stats, source = load_stats_file(local_file), f"파일: {local_file.name}"
    elif searchad_client is not None or settings.naver_searchad_ready():
        if searchad_client is None:
            from flying_geese.stage2_blue_ocean.naver_searchad import NaverSearchAdClient

            searchad_client = NaverSearchAdClient(settings)
        stats, source = fetch_stats_from_api(searchad_client, hints), f"검색광고 API (힌트 {len(hints)}개)"
    else:
        example = marketing_dir / "keyword_stats.example.json"
        stats, source, using_examples = load_stats_file(example), "⚠ 예시 데이터", True

    keywords = classify_all(stats, rules)

    # 2) (선택) 포화도: 쇼핑 상품 수
    counts: dict[str, int] = {}
    counts_file = marketing_dir / "shopping_counts.json"
    if counts_file.exists():
        counts.update(json.loads(counts_file.read_text(encoding="utf-8")))
    if with_shopping and (shopping_client is not None or settings.naver_datalab_ready()):
        if shopping_client is None:
            from flying_geese.marketing.naver_shopping import NaverShoppingClient

            shopping_client = NaverShoppingClient(settings)
        targets = [k.stat.keyword for k in keywords if k.grade in (GRADE_CORE, GRADE_TEST) and k.stat.keyword not in counts]
        counts.update(shopping_client.product_counts(targets[:SHOPPING_LOOKUP_LIMIT]))
        if not using_examples:
            counts_file.write_text(json.dumps(counts, ensure_ascii=False, indent=1), encoding="utf-8")
    if counts:
        for kw in keywords:
            if kw.stat.keyword in counts:
                kw.stat.shopping_products = counts[kw.stat.keyword]
        # 포화도가 점수에 반영되도록 재분류
        keywords = classify_all([k.stat for k in keywords], rules)

    # 3) DB 누적 (예시 데이터는 누적하지 않음)
    db_path = None
    if not using_examples:
        db_path = update_keyword_db(marketing_dir / "keyword_db.csv", keywords, today)
    else:
        for kw in keywords:
            kw.trend = "신규"

    # 4) 성과 보고서가 있으면 입찰 조정 제안
    perf_file = _first_existing(marketing_dir / "ad_performance.csv")
    if perf_file is None and using_examples:
        perf_file = _first_existing(marketing_dir / "ad_performance.example.csv")
    recommendations = recommend_all(load_performance(perf_file), BidRules()) if perf_file else []

    output_dir.mkdir(parents=True, exist_ok=True)
    report_path = write_keyword_report(
        keywords, recommendations, output_dir / "keyword_report.xlsx", today, source, using_examples
    )
    upload_path = write_upload_csv(keywords, output_dir / "keyword_upload.csv")
    return KeywordRunResult(keywords, source, report_path, upload_path, recommendations, db_path, using_examples)


def write_upload_csv(keywords: list[ClassifiedKeyword], path: Path) -> Path:
    """광고그룹별로 정리한 등록용 목록 (검색광고 대량관리 양식에 붙여넣기용)."""
    rows = [
        {
            "광고그룹": kw.ad_group,
            "키워드": kw.stat.keyword,
            "등급": kw.grade,
            "시작입찰가(원)": 70 if "롱테일" in kw.reason else "",
            "입찰전략": kw.bid_strategy,
        }
        for kw in keywords
        if kw.grade in (GRADE_CORE, GRADE_TEST)
    ]
    rows.sort(key=lambda r: (r["광고그룹"], r["등급"]))
    return write_csv_rows(path, ["광고그룹", "키워드", "등급", "시작입찰가(원)", "입찰전략"], rows)


def _kw_row(kw: ClassifiedKeyword) -> list:
    tags = [t for t, on in (("키즈", kw.is_kid), ("프리미엄", kw.is_premium), ("구매", kw.is_purchase), ("선물", kw.is_gift)) if on]
    return [
        kw.stat.keyword, kw.stat.total_volume, kw.stat.comp_idx, round(kw.stat.monthly_clicks, 1),
        kw.stat.shopping_products if kw.stat.shopping_products is not None else "",
        saturation_label(kw.stat.saturation), ", ".join(tags), kw.score, kw.ad_group, kw.bid_strategy,
        kw.trend, kw.reason,
    ]


_KW_HEADERS = ["키워드", "월검색량", "경쟁", "월클릭", "쇼핑상품수", "포화도", "태그", "점수", "광고그룹", "입찰전략", "추세", "판정사유"]
_KW_WIDTHS = [24, 9, 6, 7, 10, 7, 18, 7, 16, 30, 7, 40]


def write_keyword_report(
    keywords: list[ClassifiedKeyword],
    recommendations: list[BidRecommendation],
    path: Path,
    today: date,
    source: str,
    using_examples: bool,
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "요약"
    row = write_title(sheet, "네이버 검색광고 키워드 리포트", f"기준일 {today.isoformat()} · 출처 {source}")
    grade_rows = [[g, sum(1 for k in keywords if k.grade == g)] for g in GRADE_ORDER]
    row = write_table(sheet, row, ["등급", "키워드 수"], grade_rows, [16, 12]) + 2
    groups: dict[str, list[ClassifiedKeyword]] = {}
    for kw in keywords:
        if kw.grade in (GRADE_CORE, GRADE_TEST):
            groups.setdefault(kw.ad_group, []).append(kw)
    group_rows = [
        [name, len(kws), sum(k.stat.total_volume for k in kws), ", ".join(k.stat.keyword for k in kws[:5])]
        for name, kws in sorted(groups.items(), key=lambda x: -sum(k.stat.total_volume for k in x[1]))
    ]
    write_table(sheet, row, ["광고그룹(제안)", "키워드 수", "합계 검색량", "대표 키워드"], group_rows, [16, 12, 12, 70])
    sheet.freeze_panes = None
    if using_examples:
        sheet["A2"].value += " (예시 데이터 - 실제 키워드도구 결과로 교체 필요)"

    for grade, title in (
        (GRADE_CORE, "A 핵심(집행)"),
        (GRADE_TEST, "B 테스트"),
        (GRADE_CONTENT, "C 콘텐츠"),
    ):
        sheet = workbook.create_sheet(title)
        start = write_title(sheet, title)
        rows = [_kw_row(k) for k in keywords if k.grade == grade]
        write_table(sheet, start, _KW_HEADERS, rows, _KW_WIDTHS)

    sheet = workbook.create_sheet("신규·급상승")
    start = write_title(sheet, "신규·급상승 키워드", "직전 실행 대비 (keyword_db.csv 누적 기준)")
    movers = [k for k in keywords if k.grade != GRADE_EXCLUDE and k.trend in ("신규", "급상승")]
    write_table(sheet, start, _KW_HEADERS, [_kw_row(k) for k in movers], _KW_WIDTHS)

    sheet = workbook.create_sheet("제외")
    start = write_title(sheet, "제외 키워드 (상위 300)", "제외어 리스트는 data/marketing/keyword_rules.json에서 관리")
    excluded = [k for k in keywords if k.grade == GRADE_EXCLUDE][:300]
    write_table(sheet, start, ["키워드", "월검색량", "사유"], [[k.stat.keyword, k.stat.total_volume, k.reason] for k in excluded], [28, 10, 30])

    if recommendations:
        sheet = workbook.create_sheet("입찰 조정 제안")
        start = write_title(sheet, "입찰 조정 제안", "ROAS 목표 300% 기준 · 우선순위 1=즉시 처리")
        rows = [
            [r.priority, r.perf.ad_group, r.perf.keyword, r.perf.impressions, r.perf.clicks, f"{r.perf.ctr:.2%}",
             round(r.perf.cost), r.perf.conversions, round(r.perf.revenue),
             f"{r.perf.roas:.0%}" if r.perf.roas is not None else "-", r.action, r.reason]
            for r in recommendations
        ]
        write_table(
            sheet, start,
            ["우선", "광고그룹", "키워드", "노출", "클릭", "CTR", "비용", "전환", "매출", "ROAS", "제안", "사유"],
            rows, [6, 14, 22, 8, 7, 7, 10, 6, 10, 7, 16, 50],
        )
        for offset, r in enumerate(recommendations, start=1):
            cell = sheet.cell(row=start + offset, column=11)
            cell.font = Font(bold=r.priority == 1)
            if "+" in r.action:
                cell.fill = FILL_GOOD
            elif r.priority <= 2:
                cell.fill = FILL_WARN

    workbook.save(path)
    return path
