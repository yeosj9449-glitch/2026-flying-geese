"""신규 주문 -> 농가 전달용 발주서 엑셀 자동 생성."""
from __future__ import annotations

from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from flying_geese.models import OrderLine

HEADERS = ["주문번호", "품목", "품종", "수량", "단가", "수취인", "주소", "연락처"]


def generate_purchase_order_excel(
    orders: list[OrderLine], output_path: str | Path, issue_date: date | None = None
) -> Path:
    """공급처(supplier_id)별로 시트를 분리한 발주서 엑셀 파일을 생성한다."""
    issue_date = issue_date or date.today()
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    by_supplier: dict[str, list[OrderLine]] = defaultdict(list)
    for order in orders:
        by_supplier[order.supplier_id or "미지정"].append(order)

    workbook = Workbook()
    workbook.remove(workbook.active)

    for supplier_id, supplier_orders in by_supplier.items():
        sheet = workbook.create_sheet(title=_safe_sheet_name(supplier_id))
        _write_supplier_sheet(sheet, supplier_id, supplier_orders, issue_date)

    workbook.save(output_path)
    return output_path


def _write_supplier_sheet(
    sheet: Worksheet, supplier_id: str, orders: list[OrderLine], issue_date: date
) -> None:
    sheet["A1"] = f"발주서 - 공급처: {supplier_id} ({issue_date.isoformat()})"
    sheet["A1"].font = Font(bold=True, size=13)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))

    header_row = 3
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    total_amount = 0
    for row_idx, order in enumerate(orders, start=header_row + 1):
        amount = order.quantity * order.unit_price
        total_amount += amount
        sheet.cell(row=row_idx, column=1, value=order.order_id)
        sheet.cell(row=row_idx, column=2, value=order.product_name)
        sheet.cell(row=row_idx, column=3, value=order.variety_keyword)
        sheet.cell(row=row_idx, column=4, value=order.quantity)
        sheet.cell(row=row_idx, column=5, value=order.unit_price)
        sheet.cell(row=row_idx, column=6, value=order.buyer_name)
        sheet.cell(row=row_idx, column=7, value=order.buyer_address)
        sheet.cell(row=row_idx, column=8, value=order.buyer_phone)

    total_row = header_row + len(orders) + 1
    sheet.cell(row=total_row, column=4, value="합계 금액").font = Font(bold=True)
    sheet.cell(row=total_row, column=5, value=total_amount).font = Font(bold=True)

    for col_idx, width in enumerate([14, 16, 20, 8, 10, 10, 30, 16], start=1):
        sheet.column_dimensions[chr(64 + col_idx)].width = width


def _safe_sheet_name(name: str) -> str:
    invalid = set(r"\/*?:[]")
    cleaned = "".join(c for c in name if c not in invalid)
    return cleaned[:31] or "공급처"
