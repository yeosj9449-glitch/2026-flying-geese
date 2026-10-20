"""소싱 관리 대시보드 엑셀 생성.

시트 구성: 요약 → 공급처 평가 → 할 일 → 컨택 이력 → 샘플 검수 → 퍼널 추이 → 입점 기준표
"""
from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from flying_geese.models import (
    ContactLog,
    LeadEvaluation,
    SampleInspection,
    SourcingGrade,
    SourcingStage,
)
from flying_geese.sourcing.criteria import (
    ALLOWED_CERTS,
    CERT_REVERIFY_DAYS,
    SAFETY_TEST_LABEL,
    SAFETY_TEST_VALID_DAYS,
)
from flying_geese.tabular import read_csv_rows
from flying_geese.xlsx_style import (
    FILL_BAD,
    FILL_GOOD,
    FILL_KIDS,
    FILL_WARN,
    fill_row,
    write_table,
    write_title,
)

FOLLOW_UP_DAYS = 7  # 컨택/샘플요청 단계에서 마지막 연락 후 이 기간이 지나면 재연락
_ACTIVE_STAGES = {
    SourcingStage.DISCOVERED,
    SourcingStage.CONTACTED,
    SourcingStage.SAMPLE_REQUESTED,
    SourcingStage.SAMPLE_REVIEW,
}
_GRADE_FILL = {
    SourcingGrade.KIDS: FILL_KIDS,
    SourcingGrade.PASS: FILL_GOOD,
    SourcingGrade.FIXABLE: FILL_WARN,
    SourcingGrade.REJECT: FILL_BAD,
}


@dataclass
class TodoItem:
    due: date | None
    lead_id: str
    supplier_name: str
    action: str
    source: str  # 수동 입력 / 자동 제안 / 재연락
    overdue: bool = False


@dataclass
class SourcingSummary:
    stage_counts: dict[str, int]
    grade_counts: dict[str, int]
    kids_category_items: list[str]
    todos: list[TodoItem] = field(default_factory=list)
    warnings: list[tuple[str, str]] = field(default_factory=list)  # (공급처, 경고)


def last_contact_dates(contacts: list[ContactLog]) -> dict[str, date]:
    latest: dict[str, date] = {}
    for log in contacts:
        if log.contacted_on and (log.lead_id not in latest or log.contacted_on > latest[log.lead_id]):
            latest[log.lead_id] = log.contacted_on
    return latest


def build_todos(
    evaluations: list[LeadEvaluation], contacts: list[ContactLog], today: date
) -> list[TodoItem]:
    last_contact = last_contact_dates(contacts)
    todos: list[TodoItem] = []
    for ev in evaluations:
        lead = ev.lead
        if lead.stage == SourcingStage.DROPPED:
            continue
        if lead.next_action:
            todos.append(
                TodoItem(
                    lead.next_action_date, lead.lead_id, lead.supplier_name, lead.next_action, "수동 입력",
                    overdue=bool(lead.next_action_date and lead.next_action_date < today),
                )
            )
        if lead.stage in (SourcingStage.CONTACTED, SourcingStage.SAMPLE_REQUESTED):
            last = last_contact.get(lead.lead_id)
            if last is None or (today - last).days >= FOLLOW_UP_DAYS:
                since = f"마지막 연락 {last.isoformat()}" if last else "연락 기록 없음"
                todos.append(
                    TodoItem(today, lead.lead_id, lead.supplier_name, f"재연락 ({since})", "재연락")
                )
        if lead.stage in _ACTIVE_STAGES or lead.stage == SourcingStage.ONBOARDED:
            for action in ev.suggested_actions[:3]:
                todos.append(TodoItem(None, lead.lead_id, lead.supplier_name, action, "자동 제안"))
    # 지연 → 날짜 있는 것(빠른 순) → 날짜 없는 자동 제안
    return sorted(todos, key=lambda t: (not t.overdue, t.due is None, t.due or date.max, t.lead_id))


def summarize(
    evaluations: list[LeadEvaluation], contacts: list[ContactLog], today: date
) -> SourcingSummary:
    stage_counts = Counter(ev.lead.stage.value for ev in evaluations)
    grade_counts = Counter(ev.grade.value for ev in evaluations)
    kids_items = sorted(
        {
            item
            for ev in evaluations
            if ev.grade == SourcingGrade.KIDS and ev.lead.stage != SourcingStage.DROPPED
            for item in ev.lead.items
        }
    )
    warnings = [(ev.lead.supplier_name, w) for ev in evaluations for w in ev.warnings]
    return SourcingSummary(
        stage_counts={s.value: stage_counts.get(s.value, 0) for s in SourcingStage},
        grade_counts={g.value: grade_counts.get(g.value, 0) for g in SourcingGrade},
        kids_category_items=kids_items,
        todos=build_todos(evaluations, contacts, today),
        warnings=warnings,
    )


def _ox(flag: bool) -> str:
    return "O" if flag else "X"


def generate_sourcing_dashboard(
    evaluations: list[LeadEvaluation],
    contacts: list[ContactLog],
    samples: list[SampleInspection],
    output_path: str | Path,
    today: date,
    history_path: Path | None = None,
    using_examples: bool = False,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    summary = summarize(evaluations, contacts, today)
    names = {ev.lead.lead_id: ev.lead.supplier_name for ev in evaluations}

    workbook = Workbook()

    # --- 요약 ---
    sheet = workbook.active
    sheet.title = "요약"
    subtitle = f"기준일 {today.isoformat()} · 공급처 {len(evaluations)}곳"
    if using_examples:
        subtitle += " · ⚠ 예시 데이터(data/sourcing/*.example.csv) 기반"
    row = write_title(sheet, "소싱 관리 대시보드", subtitle)
    row = write_table(sheet, row, ["단계", "공급처 수"], [[k, v] for k, v in summary.stage_counts.items()], [16, 12]) + 2
    row = write_table(sheet, row, ["등급", "공급처 수"], [[k, v] for k, v in summary.grade_counts.items()]) + 2
    sheet.cell(row=row, column=1, value="키즈 전용 카테고리 확보 품목").font = Font(bold=True)
    sheet.cell(row=row, column=2, value=", ".join(summary.kids_category_items) or "아직 없음")
    row += 2
    overdue = [t for t in summary.todos if t.overdue]
    sheet.cell(row=row, column=1, value="지연된 후속조치").font = Font(bold=True)
    sheet.cell(row=row, column=2, value=len(overdue))
    sheet.cell(row=row + 1, column=1, value="경고").font = Font(bold=True)
    sheet.cell(row=row + 1, column=2, value=len(summary.warnings))
    sheet.column_dimensions["B"].width = 60
    sheet.freeze_panes = None

    # --- 공급처 평가 ---
    sheet = workbook.create_sheet("공급처 평가")
    headers = [
        "ID", "공급처명", "품목군", "취급품목", "지역", "단계", "등급", "점수",
        "①인증", "②검사", "③아이기준", "샘플", "미충족", "경고", "다음 제안 액션", "연락처",
    ]
    rows = []
    for ev in evaluations:
        lead = ev.lead
        sample = {True: "통과", False: "불합격", None: "-"}[ev.sample_passed]
        rows.append([
            lead.lead_id, lead.supplier_name, lead.product_type, ", ".join(lead.items), lead.region,
            lead.stage.value, ev.grade.value, ev.score, _ox(ev.cert_ok), _ox(ev.safety_ok), _ox(ev.kids_ok),
            sample, " / ".join(ev.missing), " / ".join(ev.warnings),
            ev.suggested_actions[0] if ev.suggested_actions else "", lead.phone,
        ])
    start = write_title(sheet, "공급처 평가", "등급: S 키즈전용 > A 입점가능 > B 보완필요 > X 입점불가 (①②는 필수)")
    write_table(sheet, start, headers, rows, [7, 18, 7, 18, 10, 9, 12, 6, 6, 6, 8, 7, 30, 30, 36, 14])
    for offset, ev in enumerate(evaluations, start=1):
        sheet.cell(row=start + offset, column=7).fill = _GRADE_FILL[ev.grade]
        if ev.warnings:
            sheet.cell(row=start + offset, column=14).fill = FILL_WARN

    # --- 할 일 ---
    sheet = workbook.create_sheet("할 일")
    start = write_title(sheet, "할 일", "지연 → 일정 순 → 자동 제안 순. 처리하면 leads.csv의 다음액션/단계를 갱신하세요.")
    todo_rows = [
        [t.due.isoformat() if t.due else "", "지연" if t.overdue else "", t.lead_id, t.supplier_name, t.action, t.source]
        for t in summary.todos
    ]
    write_table(sheet, start, ["기한", "상태", "ID", "공급처명", "할 일", "구분"], todo_rows, [12, 6, 7, 18, 60, 10])
    for offset, todo in enumerate(summary.todos, start=1):
        if todo.overdue:
            fill_row(sheet, start + offset, 6, FILL_BAD)

    # --- 컨택 이력 ---
    sheet = workbook.create_sheet("컨택 이력")
    start = write_title(sheet, "컨택 이력 (누적)")
    contact_rows = [
        [c.contacted_on.isoformat() if c.contacted_on else "", c.lead_id, names.get(c.lead_id, "?"), c.channel, c.summary, c.outcome]
        for c in sorted(contacts, key=lambda c: c.contacted_on or date.min, reverse=True)
    ]
    write_table(sheet, start, ["일자", "ID", "공급처명", "채널", "내용", "결과"], contact_rows, [12, 7, 18, 8, 50, 20])

    # --- 샘플 검수 ---
    sheet = workbook.create_sheet("샘플 검수")
    start = write_title(sheet, "샘플 검수", "통과 = 4개 항목 평균 4.0 이상 & 2점 이하 항목 없음")
    sample_rows = []
    for s in samples:
        scores = s.scores()
        avg = round(sum(scores) / len(scores), 2) if scores else None
        sample_rows.append([
            s.received_on.isoformat() if s.received_on else "", s.lead_id, names.get(s.lead_id, "?"), s.item,
            s.shipping_days, s.brix, s.freshness_score, s.taste_score, s.packaging_score, s.kid_fit_score, avg,
            {True: "통과", False: "불합격", None: "미입력"}[s.passed], s.notes,
        ])
    write_table(
        sheet, start,
        ["수령일", "ID", "공급처명", "품목", "배송일수", "당도", "신선도", "맛", "포장", "아이적합", "평균", "판정", "메모"],
        sample_rows, [12, 7, 18, 14, 8, 6, 7, 5, 5, 8, 6, 8, 40],
    )
    for offset, s in enumerate(samples, start=1):
        if s.passed is not None:
            sheet.cell(row=start + offset, column=12).fill = FILL_GOOD if s.passed else FILL_BAD

    # --- 퍼널 추이 ---
    if history_path and history_path.exists():
        sheet = workbook.create_sheet("퍼널 추이")
        history = read_csv_rows(history_path)
        if history:
            headers = list(history[0].keys())
            start = write_title(sheet, "퍼널 추이 (대시보드 생성 시마다 자동 누적)")
            write_table(sheet, start, headers, [[r.get(h) for h in headers] for r in history])

    # --- 기준표 ---
    sheet = workbook.create_sheet("입점 기준표")
    start = write_title(sheet, "큐레이션 입점 기준표")
    criteria_rows = [
        ["① 인증", "필수", "; ".join(f"{k}: {'/'.join(sorted(v))}" for k, v in ALLOWED_CERTS.items()),
         f"인증번호를 enviagro.go.kr에서 직접 조회, {CERT_REVERIFY_DAYS}일마다 재조회, 유효기간 내"],
        ["② 안전성 검사", "필수", "; ".join(f"{k}: {v}" for k, v in SAFETY_TEST_LABEL.items()),
         f"검사일 {SAFETY_TEST_VALID_DAYS}일 이내 성적서 보관"],
        ["③ 아이 기준", "키즈 전용 카테고리 조건", "한입 크기(소과) 선별 / 1kg 이하 소포장 / 무첨가·무염", "3개 모두 확인 시 S 등급"],
        ["④ 품질", "가점", "당도 기록, 수확→출고 2일 이내, 콜드체인", "최대 10점"],
        ["⑤ 스토리", "가점", "농가 사진·인터뷰 제공", "6점"],
        ["샘플 검수", "입점 전 필수", "신선도/맛/포장/아이적합 1~5점", "평균 4.0↑ & 2점 이하 없음"],
    ]
    write_table(sheet, start, ["항목", "구분", "내용", "판정"], criteria_rows, [14, 20, 70, 40])

    workbook.save(output_path)
    return output_path
