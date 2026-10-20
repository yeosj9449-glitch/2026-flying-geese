from datetime import date

from openpyxl import load_workbook

from flying_geese.config import Settings
from flying_geese.marketing.ad_performance import BidRules, KeywordPerformance, load_performance, recommend
from flying_geese.marketing.keyword_discovery import (
    DB_HEADERS,
    DEFAULT_MARKETING_DIR,
    GRADE_CONTENT,
    GRADE_CORE,
    GRADE_EXCLUDE,
    GRADE_TEST,
    KeywordRules,
    KeywordStat,
    build_hint_keywords,
    chunk,
    classify,
    stat_from_api_row,
    update_keyword_db,
)
from flying_geese.marketing.keyword_runner import run_keyword_discovery
from flying_geese.marketing.naver_shopping import NaverShoppingClient
from flying_geese.tabular import read_csv_rows, write_csv_rows

RULES = KeywordRules.load(DEFAULT_MARKETING_DIR / "keyword_rules.json")
NO_KEYS = Settings(
    naver_searchad_api_key="", naver_searchad_secret_key="", naver_searchad_customer_id="",
    naver_datalab_client_id="", naver_datalab_client_secret="",
)


def _stat(keyword, volume=500, comp="낮음", products=None):
    return KeywordStat(keyword, volume // 5, volume - volume // 5, comp, shopping_products=products)


def test_hint_keywords_have_no_spaces_and_chunk_by_five():
    hints = build_hint_keywords(RULES, ["감귤", "딸기"])
    assert all(" " not in h for h in hints)
    assert "유기농감귤" in hints and "이유식딸기" in hints
    assert all(len(group) <= 5 for group in chunk(hints))


def test_api_row_parsing_handles_less_than_ten_and_null():
    stat = stat_from_api_row({"relKeyword": "아기귤", "monthlyPcQcCnt": "< 10", "monthlyMobileQcCnt": 300, "compIdx": None})
    assert stat.total_volume == 300 and stat.comp_idx == ""


def test_classification_grades():
    assert classify(_stat("아기 사과"), RULES).grade == GRADE_CORE
    assert classify(_stat("사과 이유식 만드는법"), RULES).grade == GRADE_CONTENT
    assert classify(_stat("아기과자"), RULES).grade == GRADE_EXCLUDE
    assert classify(_stat("아기김치"), RULES).grade == GRADE_EXCLUDE  # '김' 품목 오탐 방지
    assert classify(_stat("유기농딸기", volume=10), RULES).grade == GRADE_TEST  # 롱테일
    assert classify(_stat("과일선물세트", volume=70000, comp="높음"), RULES).grade == GRADE_TEST
    assert classify(_stat("사과", volume=150000, comp="높음"), RULES).grade == GRADE_EXCLUDE


def test_gift_keywords_go_to_gift_ad_group_and_saturation_lowers_score():
    gift = classify(_stat("돌답례품 과일"), RULES)
    assert gift.ad_group == "선물세트"
    free = classify(_stat("아기귤", products=500), RULES)
    crowded = classify(_stat("아기귤", products=500 * 200), RULES)
    assert crowded.score < free.score


def test_keyword_db_tracks_new_and_surge_and_keeps_operator_memo(tmp_path):
    db = tmp_path / "db.csv"
    update_keyword_db(db, [classify(_stat("아기귤", 500), RULES)], date(2026, 10, 1))
    rows = read_csv_rows(db)
    rows[0]["운영상태"] = "등록완료"  # 운영자가 엑셀에서 직접 적은 값
    write_csv_rows(db, DB_HEADERS, rows)
    surged = classify(_stat("아기귤", 900), RULES)
    update_keyword_db(db, [surged], date(2026, 10, 8))
    assert surged.trend == "급상승" and surged.prev_volume == 500
    assert "등록완료" in db.read_text(encoding="utf-8-sig")


def test_bid_recommendations():
    rules = BidRules(target_roas=3.0, target_aov=60000)
    wasted = KeywordPerformance("a", "g", 5000, 100, 30000, 0, 0)
    assert recommend(wasted, rules).action == "일시중지"
    star = KeywordPerformance("b", "g", 5000, 100, 20000, 3, 150000)
    assert recommend(star, rules).action == "입찰 +15%"
    thin = KeywordPerformance("c", "g", 50, 3, 300, 0, 0)
    assert recommend(thin, rules).action == "입찰 +20%"


def test_performance_csv_with_quoted_comma_header(tmp_path):
    path = tmp_path / "p.csv"
    path.write_text('키워드,노출수,클릭수,"총비용(VAT포함,원)",전환수,전환매출액(원)\n아기귤,"1,000",50,"12,000",1,60000\n', encoding="utf-8")
    perf = load_performance(path)[0]
    assert perf.impressions == 1000 and perf.cost == 12000 and perf.roas == 5.0


class _FakeSearchAd:
    def __init__(self):
        self.calls = []

    def keyword_stats(self, hints):
        self.calls.append(hints)
        if len(self.calls) == 1:
            raise RuntimeError("일시 오류")
        return [{"relKeyword": f"아기{hints[0][-2:]}", "monthlyPcQcCnt": 100, "monthlyMobileQcCnt": 400, "compIdx": "낮음"}]


class _FakeShopping:
    def product_counts(self, keywords):
        return {k: 1000 for k in keywords}


def test_runner_with_api_survives_failed_chunk_and_writes_outputs(tmp_path):
    marketing = tmp_path / "marketing"
    marketing.mkdir()
    (marketing / "keyword_rules.json").write_text((DEFAULT_MARKETING_DIR / "keyword_rules.json").read_text(encoding="utf-8"), encoding="utf-8")
    fake = _FakeSearchAd()
    result = run_keyword_discovery(
        marketing_dir=marketing, items=["감귤", "딸기"], settings=NO_KEYS, output_dir=tmp_path / "out",
        today=date(2026, 10, 20), searchad_client=fake, with_shopping=True, shopping_client=_FakeShopping(),
    )
    assert len(fake.calls) == 2  # 힌트 10개 → 2묶음, 첫 묶음 실패해도 계속
    assert result.keywords and result.keywords[0].stat.shopping_products == 1000
    assert (marketing / "keyword_db.csv").exists() and (marketing / "shopping_counts.json").exists()
    assert "A 핵심(집행)" in load_workbook(result.report_path).sheetnames


def test_runner_falls_back_to_examples_without_keys(tmp_path):
    result = run_keyword_discovery(settings=NO_KEYS, output_dir=tmp_path, today=date(2026, 10, 20))
    assert result.using_examples and result.db_path is None
    assert result.bid_recommendations
    assert result.upload_path.read_text(encoding="utf-8-sig").startswith("광고그룹")


def test_shopping_total_parsing():
    assert NaverShoppingClient.parse_total({"total": 12345}) == 12345
    assert NaverShoppingClient.parse_total({"total": None}) == 0
