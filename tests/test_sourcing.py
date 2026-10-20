from datetime import date

from openpyxl import load_workbook

from flying_geese.models import SampleInspection, SourcingGrade, SourcingStage, SupplierLead
from flying_geese.sourcing import store
from flying_geese.sourcing.criteria import evaluate_lead, evaluate_leads
from flying_geese.sourcing.runner import log_contact, run_sourcing_dashboard

TODAY = date(2026, 10, 20)


def _lead(**overrides) -> SupplierLead:
    base = dict(
        lead_id="L1",
        supplier_name="테스트농원",
        product_type="과일",
        items=["감귤"],
        cert_type="유기농",
        cert_number="12-345",
        cert_expiry=date(2027, 6, 30),
        cert_verified_on=date(2026, 10, 1),
        safety_test_on=date(2026, 9, 1),
        kid_bite_size=True,
        kid_small_pack=True,
        kid_no_additive=True,
    )
    base.update(overrides)
    return SupplierLead(**base)


def test_all_mandatory_and_kid_checks_give_kids_grade():
    ev = evaluate_lead(_lead(), TODAY)
    assert ev.grade == SourcingGrade.KIDS
    assert ev.cert_ok and ev.safety_ok and ev.kids_ok


def test_kid_checks_missing_still_pass_but_not_kids_category():
    ev = evaluate_lead(_lead(kid_small_pack=False), TODAY)
    assert ev.grade == SourcingGrade.PASS
    assert any("소포장" in a for a in ev.suggested_actions)


def test_gap_only_is_rejected_regardless_of_score():
    """GAP은 친환경 인증이 아니므로 다른 항목이 모두 좋아도 입점 불가."""
    ev = evaluate_lead(_lead(cert_type="GAP", story_available=True, cold_chain=True, brix=15), TODAY)
    assert ev.grade == SourcingGrade.REJECT
    assert not ev.cert_ok


def test_unverified_cert_number_is_fixable_not_pass():
    ev = evaluate_lead(_lead(cert_verified_on=None), TODAY)
    assert ev.grade == SourcingGrade.FIXABLE
    assert any("enviagro" in a for a in ev.suggested_actions)


def test_cert_verification_older_than_90_days_requires_reverify():
    ev = evaluate_lead(_lead(cert_verified_on=date(2026, 6, 1)), TODAY)
    assert not ev.cert_ok
    assert any("재조회" in a for a in ev.suggested_actions)


def test_expired_cert_and_old_safety_test_fail():
    ev = evaluate_lead(_lead(cert_expiry=date(2026, 10, 1), safety_test_on=date(2025, 9, 1)), TODAY)
    assert not ev.cert_ok and not ev.safety_ok
    assert ev.grade == SourcingGrade.FIXABLE


def test_cert_aliases_and_seafood_certs():
    assert evaluate_lead(_lead(cert_type="유기농산물 인증"), TODAY).cert_ok
    assert evaluate_lead(_lead(product_type="수산", cert_type="msc", cold_chain=True), TODAY).cert_ok
    assert not evaluate_lead(_lead(product_type="수산", cert_type="유기농"), TODAY).cert_ok


def test_onboarded_without_sample_pass_warns():
    ev = evaluate_lead(_lead(stage=SourcingStage.ONBOARDED), TODAY)
    assert any("샘플" in w for w in ev.warnings)


def test_sample_pass_rule():
    assert SampleInspection("L1", "감귤", TODAY, freshness_score=5, taste_score=4, packaging_score=4, kid_fit_score=4).passed
    # 평균은 4 이상이지만 2점 항목이 있으면 불합격
    assert not SampleInspection("L1", "감귤", TODAY, freshness_score=5, taste_score=5, packaging_score=2, kid_fit_score=5).passed
    assert SampleInspection("L1", "감귤", TODAY).passed is None


def test_passed_sample_in_review_suggests_onboarding():
    sample = SampleInspection("L1", "감귤", TODAY, freshness_score=5, taste_score=5, packaging_score=4, kid_fit_score=5)
    ev = evaluate_lead(_lead(stage=SourcingStage.SAMPLE_REVIEW), TODAY, [sample])
    assert ev.sample_passed is True
    assert ev.suggested_actions[0].startswith("입점 확정")


def test_evaluate_leads_sorts_by_grade_then_score():
    leads = [_lead(lead_id="X", cert_type="GAP"), _lead(lead_id="A", kid_bite_size=False), _lead(lead_id="S")]
    assert [e.lead.lead_id for e in evaluate_leads(leads, TODAY)] == ["S", "A", "X"]


def test_csv_roundtrip_tolerates_messy_values(tmp_path):
    path = tmp_path / "leads.csv"
    path.write_text(
        "ID,공급처명,품목군,취급품목,단계,인증종류,인증번호,인증조회일,검사일,한입크기,소포장,무첨가,당도Brix\n"
        "L9,농원,과일,\"사과; 배\",없는단계,유기농,1-1,2026.10.01,2026/09/01,예,O,y,\"1,2\"\n",
        encoding="utf-8",
    )
    lead = store.load_leads(path)[0]
    assert lead.items == ["사과", "배"]
    assert lead.stage == SourcingStage.DISCOVERED
    assert lead.cert_verified_on == date(2026, 10, 1)
    assert lead.kid_bite_size and lead.kid_small_pack and lead.kid_no_additive


def test_dashboard_from_examples_does_not_write_history(tmp_path):
    report = run_sourcing_dashboard(output_path=tmp_path / "d.xlsx", today=TODAY)
    assert report.using_examples
    assert not (store.DEFAULT_SOURCING_DIR / "history.csv").exists()
    sheets = load_workbook(report.dashboard_path).sheetnames
    assert sheets[:5] == ["요약", "공급처 평가", "할 일", "컨택 이력", "샘플 검수"]
    assert any(t.overdue for t in report.summary.todos)


def test_real_data_accumulates_history_and_contacts(tmp_path):
    (tmp_path / "leads.csv").write_text(
        (store.DEFAULT_SOURCING_DIR / "leads.example.csv").read_text(encoding="utf-8"), encoding="utf-8"
    )
    log_contact("L003", "전화", "검사성적서 요청", stage="샘플요청", next_action="샘플 수령", sourcing_dir=tmp_path, today=TODAY)
    run_sourcing_dashboard(sourcing_dir=tmp_path, output_path=tmp_path / "d.xlsx", today=TODAY)
    run_sourcing_dashboard(sourcing_dir=tmp_path, output_path=tmp_path / "d.xlsx", today=date(2026, 10, 27))

    assert len(store.load_contacts(tmp_path / "contacts.csv")) == 1
    updated = {l.lead_id: l for l in store.load_leads(tmp_path / "leads.csv")}["L003"]
    assert updated.stage == SourcingStage.SAMPLE_REQUESTED and updated.next_action == "샘플 수령"
    history = (tmp_path / "history.csv").read_text(encoding="utf-8-sig").splitlines()
    assert len(history) == 3  # 헤더 + 2일치


def test_log_contact_rejects_unknown_id(tmp_path):
    (tmp_path / "leads.csv").write_text("ID,공급처명\nL1,농원\n", encoding="utf-8")
    try:
        log_contact("NOPE", "전화", "x", sourcing_dir=tmp_path)
    except ValueError:
        return
    raise AssertionError("ValueError expected")


def test_unknown_cert_is_fixable_with_confirmation_action():
    """발굴 단계에서 인증 정보를 아직 모르는 리드는 탈락이 아니라 확인 대상."""
    ev = evaluate_lead(_lead(cert_type="", cert_number=""), TODAY)
    assert ev.grade == SourcingGrade.FIXABLE
    assert "인증 미확인" in ev.missing
    assert not any("탈락" in a for a in ev.suggested_actions)
