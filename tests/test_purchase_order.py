from openpyxl import load_workbook

from flying_geese.models import OrderLine
from flying_geese.stage5_automation.purchase_order import generate_purchase_order_excel


def test_generate_purchase_order_excel_groups_by_supplier(tmp_path):
    orders = [
        OrderLine("ORD-1", "고구마", "베니하르카 호박고구마", 2, 32000, "S1", "김민준", "서울시", "010-0000-0000"),
        OrderLine("ORD-2", "사과", "감홍사과", 1, 59000, "S2", "이서연", "성남시", "010-1111-1111"),
        OrderLine("ORD-3", "고구마", "베니하르카 호박고구마", 3, 32000, "S1", "박지호", "부산시", "010-2222-2222"),
    ]
    output_path = tmp_path / "po.xlsx"

    result_path = generate_purchase_order_excel(orders, output_path)

    assert result_path.exists()
    workbook = load_workbook(result_path)
    assert set(workbook.sheetnames) == {"S1", "S2"}

    s1_rows = list(workbook["S1"].iter_rows(values_only=True))
    order_ids = [row[0] for row in s1_rows if row[0] and str(row[0]).startswith("ORD")]
    assert order_ids == ["ORD-1", "ORD-3"]

    total_row = s1_rows[-1]
    assert total_row[4] == 2 * 32000 + 3 * 32000
