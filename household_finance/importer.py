"""자동화 도구: 카드·은행 내역 가져오기(자동 분류·중복 제거)와 월말 스냅샷 생성.

카드사·은행 홈페이지에서 내려받은 내역을 CSV(또는 .xlsx)로 저장해 넣으면 된다.
열 이름은 회사마다 달라서 흔히 쓰는 이름(이용일자/가맹점명/이용금액, 거래일자/적요/출금액 등)을
자동으로 찾는다. 찾지 못하면 --date-col 등으로 직접 지정한다.
"""

from __future__ import annotations

import csv
import io
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from household_finance import schema as s
from household_finance.workbook import F_INPUT, MONEY, add_charts, put, write_tx_row

DATE_KEYS = ["이용일자", "이용일", "거래일자", "거래일", "승인일자", "승인일", "거래일시", "일자", "날짜", "date"]
DESC_KEYS = ["가맹점명", "이용가맹점", "가맹점", "이용하신곳", "적요", "거래내용", "기재내용", "내용", "사용처", "description"]
AMOUNT_KEYS = ["이용금액", "결제금액", "승인금액", "거래금액", "금액", "amount"]
OUT_KEYS = ["출금액", "출금금액", "찾으신금액", "출금"]
IN_KEYS = ["입금액", "입금금액", "맡기신금액", "입금"]


@dataclass
class StatementRow:
    day: date
    desc: str
    amount: int
    is_income: bool = False


@dataclass
class ImportedTx:
    day: date
    desc: str
    amount: int
    category: str
    member: str
    payment: str
    source: str
    memo: str = ""


@dataclass
class ImportResult:
    added: list[ImportedTx] = field(default_factory=list)
    duplicates: int = 0
    skipped: list[str] = field(default_factory=list)

    @property
    def uncategorized(self) -> list[ImportedTx]:
        return [t for t in self.added if t.category == s.UNCATEGORIZED]


# ---------------------------------------------------------------- 분류
def load_rules(wb) -> list[tuple[str, str]]:
    ws = wb[s.RULES]
    rules = []
    for kw, cat, *_ in ws.iter_rows(min_row=2, values_only=True):
        if kw and cat:
            rules.append((str(kw).strip(), str(cat).strip()))
    return rules


def categorize(desc: str, rules: list[tuple[str, str]]) -> str:
    """위에서부터 먼저 맞는 키워드 규칙의 카테고리. 없으면 '미분류'."""
    text = desc.lower()
    for keyword, category in rules:
        if keyword.lower() in text:
            return category
    return s.UNCATEGORIZED


# ---------------------------------------------------------------- 내역 파일 읽기
def _parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    text = str(value).strip()
    head = text[:10]
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(head, fmt).date()
        except ValueError:
            pass
    digits = "".join(ch for ch in text if ch.isdigit())
    if len(digits) >= 8:
        try:
            return datetime.strptime(digits[:8], "%Y%m%d").date()
        except ValueError:
            return None
    return None


def _parse_amount(value) -> int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(round(value))
    text = str(value).replace(",", "").replace("원", "").replace(" ", "").strip()
    if text in ("", "-"):
        return None
    try:
        return int(round(float(text)))
    except ValueError:
        return None


def _read_table(path: Path) -> list[list]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(path, data_only=True, read_only=True)
        return [list(row) for row in wb.worksheets[0].iter_rows(values_only=True)]
    if path.suffix.lower() == ".xls":
        raise ValueError("구형 .xls 파일은 지원하지 않습니다. 엑셀에서 'CSV UTF-8' 또는 .xlsx로 다시 저장해 주세요.")
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp949"):   # 국내 금융사 CSV는 대부분 cp949(EUC-KR)
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"파일 인코딩을 알 수 없습니다: {path}")
    return [row for row in csv.reader(io.StringIO(text))]


def _find(header: list[str], keys: list[str], override: str | None) -> int | None:
    names = [str(h).strip() if h is not None else "" for h in header]
    if override:
        return names.index(override) if override in names else None
    for key in keys:   # 우선순위가 높은 이름부터 정확히 일치
        if key in names:
            return names.index(key)
    for key in keys:   # 다음으로 부분 일치 (예: '이용금액(원)')
        for i, name in enumerate(names):
            if key in name:
                return i
    return None


def read_statement(path: str | Path, date_col: str | None = None, desc_col: str | None = None,
                   amount_col: str | None = None) -> tuple[list[StatementRow], list[str]]:
    """카드·은행 내역 파일 → (거래 목록, 건너뛴 행 설명). 머리말이 있는 파일도 헤더 행을 찾아낸다."""
    rows = _read_table(Path(path))
    header_index = None
    for i, row in enumerate(rows[:30]):
        if _find(row, DATE_KEYS, date_col) is not None and _find(row, DESC_KEYS, desc_col) is not None:
            header_index = i
            break
    if header_index is None:
        raise ValueError("헤더 행(날짜·내용 열)을 찾지 못했습니다. --date-col/--desc-col/--amount-col로 열 이름을 지정하세요.")
    header = rows[header_index]
    i_date, i_desc = _find(header, DATE_KEYS, date_col), _find(header, DESC_KEYS, desc_col)
    i_out, i_in = _find(header, OUT_KEYS, None), _find(header, IN_KEYS, None)
    # 은행 내역처럼 입금/출금 열이 따로 있으면 그쪽을 쓴다 ('금액' 부분 일치가 '출금액'을 잡지 않도록).
    i_amount = _find(header, AMOUNT_KEYS, amount_col) if amount_col or (i_out is None and i_in is None) else None
    if i_amount is None and i_out is None and i_in is None:
        raise ValueError("금액 열을 찾지 못했습니다. --amount-col로 지정하세요.")

    parsed, skipped = [], []
    for line_no, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        cell = lambda i: row[i] if i is not None and i < len(row) else None
        day = _parse_date(cell(i_date))
        desc = str(cell(i_desc) or "").strip()
        if day is None:
            if any(c not in (None, "") for c in row):
                skipped.append(f"{line_no}행: 날짜 없음/합계 행")
            continue
        if i_amount is not None and _parse_amount(cell(i_amount)) is not None:
            amount, is_income = _parse_amount(cell(i_amount)), False
        else:
            out_amt, in_amt = _parse_amount(cell(i_out)) or 0, _parse_amount(cell(i_in)) or 0
            amount, is_income = (in_amt, True) if in_amt > 0 else (out_amt, False)
        if amount is None or amount <= 0:
            skipped.append(f"{line_no}행: 금액 0 또는 취소·환불({desc})")
            continue
        parsed.append(StatementRow(day, desc, amount, is_income))
    return parsed, skipped


# ---------------------------------------------------------------- 워크북 쓰기
def _last_data_row(ws, key_col: int = 1) -> int:
    last = 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=key_col).value not in (None, ""):
            last = r
    return last


def _tx_key(day, desc, amount, member) -> tuple:
    if isinstance(day, datetime):
        day = day.date()
    return (day, str(desc).strip(), int(amount or 0), str(member).strip())


def _save(wb, path: Path) -> None:
    add_charts(wb)                       # openpyxl은 기존 차트를 보존하지 못하므로 다시 그린다
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def plan_import(statement: list[StatementRow], existing: list[tuple], rules: list[tuple[str, str]], member: str,
                payment: str, source: str) -> ImportResult:
    """가져올 거래 결정: 자동 분류 + 중복 제거. existing은 (일자, 내용, 금액, 구성원) 목록.
    같은 거래는 개수까지 비교한다 - 이미 있는 만큼은 건너뛰고 초과분만 추가 (같은 날 커피 두 잔)."""
    if member not in s.OWNERS:
        raise ValueError(f"구성원은 {', '.join(s.OWNERS)} 중 하나여야 합니다: {member}")
    counts = Counter(_tx_key(*row) for row in existing)
    result = ImportResult()
    for row in statement:
        key = _tx_key(row.day, row.desc, row.amount, member)
        if counts[key] > 0:
            counts[key] -= 1
            result.duplicates += 1
            continue
        category = categorize(row.desc, rules)
        memo = "입금" if row.is_income and category == s.UNCATEGORIZED else ""
        result.added.append(ImportedTx(row.day, row.desc, row.amount, category, member, payment, source, memo))
    return result


def import_statement(workbook_path: str | Path, statement_path: str | Path, member: str, payment: str = "",
                     dry_run: bool = False, **columns) -> ImportResult:
    """내역 파일을 거래내역 시트에 추가. 같은 (일자, 내용, 금액, 구성원) 거래는 개수까지 비교해 중복 제거."""
    if member not in s.OWNERS:
        raise ValueError(f"구성원은 {', '.join(s.OWNERS)} 중 하나여야 합니다: {member}")
    workbook_path = Path(workbook_path)
    wb = load_workbook(workbook_path)
    ws = wb[s.TX]
    rules = load_rules(wb)
    statement, skipped = read_statement(statement_path, **columns)

    last = _last_data_row(ws)
    existing = [(ws.cell(r, 1).value, ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 8).value)
                for r in range(2, last + 1) if ws.cell(r, 1).value]
    result = plan_import(statement, existing, rules, member, payment, Path(statement_path).name)
    result.skipped = skipped

    if dry_run or not result.added:
        return result
    for offset, t in enumerate(sorted(result.added, key=lambda t: t.day), start=1):
        write_tx_row(ws, last + offset, t)
    ws.auto_filter.ref = f"A1:K{last + len(result.added)}"
    _save(wb, workbook_path)
    return result


def _shift_month(month: str, k: int) -> str:
    y, m = int(month[:4]), int(month[5:7])
    total = y * 12 + (m - 1) + k
    return f"{total // 12}-{total % 12 + 1:02d}"


def new_month(workbook_path: str | Path, month: str | None = None) -> tuple[str, int]:
    """다음 달 스냅샷 행을 만든다: 직전 달 잔액을 복사하고(사용자는 바뀐 잔액만 고침),
    새로 등록된 계좌는 0원으로 추가하며, 대시보드 기준월을 새 달로 바꾼다."""
    workbook_path = Path(workbook_path)
    wb = load_workbook(workbook_path)
    ws = wb[s.SNAPSHOT]
    last = _last_data_row(ws)
    months = sorted({str(ws.cell(r, 1).value) for r in range(2, last + 1) if ws.cell(r, 1).value})
    if month is None:
        if not months:
            raise ValueError("기존 스냅샷이 없습니다. --month yyyy-mm로 첫 달을 지정하세요.")
        month = _shift_month(months[-1], 1)
    if month in months:
        raise ValueError(f"{month} 스냅샷이 이미 있습니다.")
    previous = {}
    if months:
        prev_month = max(m for m in months if m < month) if any(m < month for m in months) else None
        for r in range(2, last + 1):
            if str(ws.cell(r, 1).value) == prev_month:
                previous[ws.cell(r, 2).value] = ws.cell(r, 3).value or 0
    accounts = [row[0] for row in wb[s.ACCOUNTS].iter_rows(min_row=2, max_row=s.ACCOUNT_ROWS + 1, values_only=True) if row[0]]
    for i, acc in enumerate(accounts, start=1):
        r = last + i
        put(ws, f"A{r}", month, "@", font=F_INPUT)
        put(ws, f"B{r}", acc, font=F_INPUT)
        put(ws, f"C{r}", previous.get(acc, 0), MONEY, font=F_INPUT)
        for name, formula in s.snapshot_row_formulas(r).items():
            put(ws, f"{s.col_letter(s.SNAPSHOT_COLUMNS, name)}{r}", formula)
    ws.auto_filter.ref = f"A1:G{last + len(accounts)}"
    wb[s.DASHBOARD]["C4"] = month
    _save(wb, workbook_path)
    return month, len(accounts)
