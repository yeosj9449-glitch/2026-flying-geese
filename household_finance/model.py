"""가계 데이터 모델 - 스프레드시트·DB·웹앱이 같은 형태로 주고받는 입력 데이터.

스프레드시트의 '입력 시트'(거래내역·스냅샷·계좌 …)와 설정값을 그대로 옮긴 것이다.
보고 수치(손익·재무상태표·대시보드)는 여기에 두지 않고 engine.py가 계산한다.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from household_finance import schema as s


@dataclass
class Account:
    id: str
    name: str
    institution: str
    type: str
    owner: str
    liquidity: str
    memo: str = ""


@dataclass
class Snapshot:
    month: str
    account_id: str
    amount: int


@dataclass
class Transaction:
    day: date
    desc: str
    amount: int
    category: str
    member: str
    payment: str = ""
    source: str = ""
    memo: str = ""
    id: int | None = None

    @property
    def month(self) -> str:
        return f"{self.day.year}-{self.day.month:02d}"

    @property
    def month_no(self) -> int:
        return self.day.year * 100 + self.day.month


@dataclass
class Category:
    name: str
    group: str
    essential: str = "N"


@dataclass
class Budget:
    category: str
    basis: str      # "월" | "연"
    amount: int


@dataclass
class Holding:
    account_id: str
    name: str
    code: str
    asset_class: str
    qty: float
    avg_price: float
    price: float
    id: int | None = None


@dataclass
class Debt:
    name: str
    account_id: str
    rate: float
    maturity: date | None
    method: str
    payment: int
    id: int | None = None


@dataclass
class Goal:
    name: str
    owner: str
    target: int | None          # None = 비상자금처럼 자동 계산
    start: date
    end: date
    accounts: list[str] = field(default_factory=list)
    id: int | None = None


@dataclass
class Insurance:
    insured: str
    insurer: str
    product: str
    coverages: dict[str, int]   # 보장항목 → 보장금액
    premium: int
    maturity: date | None = None
    renewable: str = "N"
    id: int | None = None


@dataclass
class Event:
    day: date
    desc: str
    kind: str       # 유입 | 유출 | 알림
    amount: int
    member: str
    id: int | None = None


@dataclass
class Settings:
    base_month: str
    t_saving: float = 0.30
    t_emergency: float = 6
    t_debt: float = 0.40
    t_dsr: float = 0.30
    band: float = 0.05
    invest_targets: dict[str, int] = field(default_factory=lambda: {"본인": 1_150_000, "배우자": 850_000, "공동": 200_000})
    min_cash: int = 3_000_000
    alert_days: int = 30
    goal_tol: float = 0.05
    lim_pension_total: int = 9_000_000
    lim_pension: int = 6_000_000
    lim_isa: int = 20_000_000

    def invest_target(self, view: str) -> int:
        return self.invest_targets.get(view, 0) if view in s.OWNERS else sum(self.invest_targets.values())


@dataclass
class Dataset:
    settings: Settings
    categories: list[Category] = field(default_factory=list)
    rules: list[tuple[str, str]] = field(default_factory=list)
    accounts: list[Account] = field(default_factory=list)
    snapshots: list[Snapshot] = field(default_factory=list)
    transactions: list[Transaction] = field(default_factory=list)
    budgets: list[Budget] = field(default_factory=list)
    holdings: list[Holding] = field(default_factory=list)
    allocation: dict[str, float] = field(default_factory=lambda: dict(s.DEFAULT_ALLOCATION))
    debts: list[Debt] = field(default_factory=list)
    goals: list[Goal] = field(default_factory=list)
    insurances: list[Insurance] = field(default_factory=list)
    events: list[Event] = field(default_factory=list)


def default_categories() -> list[Category]:
    return [Category(name, group, ess) for name, group, ess, _b, _a in s.CATEGORY_TABLE]


def default_budgets() -> list[Budget]:
    return [Budget(name, basis, amount) for name, group, _e, basis, amount in s.CATEGORY_TABLE if group in s.CONSUMPTION_GROUPS]


def empty_dataset(base_month: str) -> Dataset:
    return Dataset(settings=Settings(base_month), categories=default_categories(), rules=list(s.DEFAULT_RULES),
                   budgets=default_budgets(),
                   goals=[Goal("비상자금", "공동", None, date.today(), date(date.today().year + 1, date.today().month, min(date.today().day, 28)))])


def sample_dataset() -> Dataset:
    """sample_data 모듈의 시연용 데이터 (스프레드시트 샘플 파일과 같은 내용)."""
    from household_finance import sample_data as sd
    return Dataset(
        settings=Settings(sd.BASE_MONTH),
        categories=default_categories(),
        rules=list(s.DEFAULT_RULES),
        accounts=[Account(*a) for a in sd.ACCOUNTS],
        snapshots=[Snapshot(m, a, v) for m, a, v in sd.snapshots()],
        transactions=[Transaction(t.day, t.desc, t.amount, t.category, t.member, t.payment, t.source) for t in sd.transactions()],
        budgets=default_budgets(),
        holdings=[Holding(*h) for h in sd.holdings()],
        debts=[Debt(*d) for d in sd.DEBTS],
        goals=[Goal(n, o, t, st, en, [a for a in (a1, a2, a3) if a]) for n, o, t, st, en, a1, a2, a3 in sd.GOALS],
        insurances=[Insurance(i[0], i[1], i[2], dict(zip(s.COVERAGES, i[3:8])), i[8], i[9], i[10]) for i in sd.INSURANCES],
        events=[Event(*e) for e in sd.EVENTS],
    )


# ---------------------------------------------------------------- 스프레드시트 → 데이터셋
def _d(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _num(v, default=0):
    if v is None or v == "" or (isinstance(v, str) and v.startswith("=")):
        return default
    try:
        return float(v) if isinstance(v, float) else int(v)
    except (TypeError, ValueError):
        return default


def _txt(v) -> str:
    return "" if v is None else str(v).strip()


def load_workbook_dataset(path: str | Path) -> Dataset:
    """1단계 스프레드시트(.xlsx)의 입력 시트를 읽어 데이터셋으로 만든다 (웹앱 이전용).

    수식 칸은 건너뛰고 사람이 입력한 값만 읽는다. 목표금액이 수식(비상자금 자동 계산)이면 None.
    """
    from openpyxl import load_workbook

    from household_finance.workbook import BUDGET_FIRST, BUDGET_LAST, DEBT_FIRST, EVENT_FIRST, GOAL_FIRST, INS_FIRST, INV

    wb = load_workbook(path)
    st = wb[s.SETTINGS]
    base = wb[s.DASHBOARD]["C4"].value
    base = base.strftime("%Y-%m") if isinstance(base, (date, datetime)) else _txt(base)
    settings = Settings(
        base_month=base,
        t_saving=_num(st["B12"].value, 0.30), t_emergency=_num(st["B13"].value, 6), t_debt=_num(st["B14"].value, 0.40),
        t_dsr=_num(st["B15"].value, 0.30), band=_num(st["B16"].value, 0.05), min_cash=_num(st["B18"].value, 3_000_000),
        alert_days=_num(st["B19"].value, 30), goal_tol=_num(st["B20"].value, 0.05),
        lim_pension_total=_num(st["B22"].value, 9_000_000), lim_pension=_num(st["B23"].value, 6_000_000),
        lim_isa=_num(st["B24"].value, 20_000_000),
        invest_targets={m: _num(st[f"B{r}"].value, 0) for m, r in zip(s.OWNERS, (26, 27, 28))},
    )
    ds = Dataset(settings=settings)
    for name, group, ess, *_ in wb[s.CATEGORIES].iter_rows(min_row=2, max_col=3, values_only=True):
        if name:
            ds.categories.append(Category(_txt(name), _txt(group), _txt(ess) or "N"))
    ds.rules = [(_txt(k), _txt(c)) for k, c, *_ in wb[s.RULES].iter_rows(min_row=2, max_col=2, values_only=True) if k and c]
    for row in wb[s.ACCOUNTS].iter_rows(min_row=2, max_row=s.ACCOUNT_ROWS + 1, max_col=7, values_only=True):
        if row[0]:
            ds.accounts.append(Account(*[_txt(v) for v in row]))
    for month, acc, amount, *_ in wb[s.SNAPSHOT].iter_rows(min_row=2, max_col=3, values_only=True):
        if month and acc:
            month = month.strftime("%Y-%m") if isinstance(month, (date, datetime)) else _txt(month)
            ds.snapshots.append(Snapshot(month, _txt(acc), _num(amount)))
    for r in wb[s.TX].iter_rows(min_row=2, max_col=11, values_only=True):
        day = _d(r[0])
        if day:
            ds.transactions.append(Transaction(day, _txt(r[2]), _num(r[3]), _txt(r[4]), _txt(r[7]), _txt(r[8]), _txt(r[9]), _txt(r[10])))
    bws = wb[s.BUDGET]
    for r in range(BUDGET_FIRST, BUDGET_LAST + 1):
        cat = _txt(bws[f"A{r}"].value)
        if cat:
            ds.budgets.append(Budget(cat, _txt(bws[f"C{r}"].value) or "월", _num(bws[f"D{r}"].value)))
    inv = wb[s.INVEST]
    ds.allocation = {_txt(inv[f"A{r}"].value): _num(inv[f"B{r}"].value, 0.0)
                     for r in range(INV["alloc_first"], INV["alloc_total"]) if inv[f"A{r}"].value}
    for r in range(INV["hold_first"], INV["hold_first"] + s.HOLDING_ROWS):
        if inv[f"A{r}"].value:
            ds.holdings.append(Holding(_txt(inv[f"A{r}"].value), _txt(inv[f"D{r}"].value), _txt(inv[f"E{r}"].value),
                                       _txt(inv[f"F{r}"].value), _num(inv[f"G{r}"].value), _num(inv[f"H{r}"].value),
                                       _num(inv[f"I{r}"].value)))
    dws = wb[s.DEBT]
    for r in range(DEBT_FIRST, DEBT_FIRST + s.DEBT_ROWS):
        if dws[f"B{r}"].value:
            ds.debts.append(Debt(_txt(dws[f"A{r}"].value), _txt(dws[f"B{r}"].value), _num(dws[f"E{r}"].value, 0.0),
                                 _d(dws[f"F{r}"].value), _txt(dws[f"G{r}"].value), _num(dws[f"H{r}"].value)))
    gws = wb[s.GOALS]
    for r in range(GOAL_FIRST, GOAL_FIRST + s.GOAL_ROWS):
        if gws[f"A{r}"].value:
            target = gws[f"C{r}"].value
            target = None if isinstance(target, str) and target.startswith("=") else _num(target)
            ds.goals.append(Goal(_txt(gws[f"A{r}"].value), _txt(gws[f"B{r}"].value), target,
                                 _d(gws[f"D{r}"].value), _d(gws[f"E{r}"].value),
                                 [_txt(gws[f"{c}{r}"].value) for c in "FGH" if gws[f"{c}{r}"].value]))
    iws = wb[s.INSURANCE]
    for r in range(INS_FIRST, INS_FIRST + s.INSURANCE_ROWS):
        if iws[f"A{r}"].value:
            cov = {c: _num(iws.cell(r, 4 + i).value) for i, c in enumerate(s.COVERAGES)}
            ds.insurances.append(Insurance(_txt(iws[f"A{r}"].value), _txt(iws[f"B{r}"].value), _txt(iws[f"C{r}"].value),
                                           cov, _num(iws[f"I{r}"].value), _d(iws[f"J{r}"].value), _txt(iws[f"K{r}"].value) or "N"))
    ews = wb[s.EVENTS]
    for r in range(EVENT_FIRST, EVENT_FIRST + s.EVENT_ROWS):
        day = _d(ews[f"A{r}"].value)
        if day:
            ds.events.append(Event(day, _txt(ews[f"B{r}"].value), _txt(ews[f"C{r}"].value), _num(ews[f"D{r}"].value),
                                   _txt(ews[f"E{r}"].value)))
    return ds
